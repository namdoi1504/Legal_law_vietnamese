from io import BytesIO
import unittest
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient
from configs.settings import settings
from app.main import create_app
from src.parser.document_reader import DocumentReader
from src.pipeline.rag_pipeline import RAGResult


class APITests(unittest.TestCase):
    def setUp(self):
        self.pipeline = Mock()
        self.pipeline.run.return_value = RAGResult(
            answer="Thiếu dữ kiện.", status="needs_information"
        )
        self.document = Mock()
        self.client = TestClient(create_app(self.pipeline, self.document))
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)

    def test_chat_preserves_contract_and_date_filter(self):
        response = self.client.post(
            "/chat", json={"query": "Câu hỏi?", "as_of_date": "2022-01-01"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("latency_ms", response.json())
        self.assertEqual(response.json()["status"], "needs_information")
        self.assertNotIn("retrieved_chunks", response.json())
        self.assertEqual(self.pipeline.run.call_args.kwargs["as_of_date"], "2022-01-01")

    def test_empty_and_oversized_queries(self):
        self.assertEqual(self.client.post("/chat", json={"query": ""}).status_code, 422)
        self.assertEqual(
            self.client.post(
                "/chat", json={"query": "x" * (settings.MAX_QUERY_CHARS + 1)}
            ).status_code,
            422,
        )

    def test_provider_error_is_503(self):
        self.pipeline.run.return_value = RAGResult(
            answer="Không kết nối được.", status="provider_error"
        )
        self.assertEqual(
            self.client.post("/chat", json={"query": "Câu hỏi?"}).status_code, 503
        )

    def test_quota_error_is_429_with_retry_header(self):
        self.pipeline.run.return_value = RAGResult(answer='Hết quota.', status='rate_limited', retry_after_seconds=12)
        response = self.client.post('/chat', json={'query':'Câu hỏi?'})
        self.assertEqual(response.status_code, 429)
        self.assertEqual(response.headers['Retry-After'], '12')

    def test_rejects_doc_and_oversized_upload(self):
        self.assertEqual(
            self.client.post(
                "/analyze-document", files={"files": ("old.doc", b"abc")}
            ).status_code,
            400,
        )
        with patch.object(settings, "MAX_UPLOAD_BYTES", 10):
            self.assertEqual(
                self.client.post(
                    "/analyze-document", files={"files": ("new.pdf", b"x" * 11)}
                ).status_code,
                413,
            )

    def test_health_does_not_call_provider(self):
        self.assertEqual(self.client.get("/health").status_code, 200)
        self.pipeline.run.assert_not_called()


class ReaderTests(unittest.TestCase):
    def test_blank_separator_page_is_not_mistaken_for_scan(self):
        import fitz
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((72,72), 'Contract text')
            doc.new_page()
            text = DocumentReader().read_bytes(doc.tobytes(), 'contract.pdf')
        self.assertIn('Contract text', text)

    def test_rejects_disguised_and_scan_files(self):
        reader = DocumentReader()
        for data, name in [
            (b"not a pdf", "x.pdf"),
            (b"not a zip", "x.docx"),
            (b"word", "x.doc"),
        ]:
            with self.assertRaises(ValueError):
                reader.read_bytes(data, name)
        import fitz

        with fitz.open() as doc:
            doc.new_page()
            with self.assertRaisesRegex(ValueError, "OCR"):
                reader.read_bytes(doc.tobytes(), "scan.pdf")

    def test_docx_keeps_table_in_document_order(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Before")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Salary clause"
        doc.add_paragraph("After")
        stream = BytesIO()
        doc.save(stream)
        text = DocumentReader().read_bytes(stream.getvalue(), "contract.docx")
        self.assertLess(text.index("Before"), text.index("Salary clause"))
        self.assertLess(text.index("Salary clause"), text.index("After"))


if __name__ == "__main__":
    unittest.main()
